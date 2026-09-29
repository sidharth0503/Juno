from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DARK = "172B42"
BLUE = "0000FF"
RED = "A32323"
NUMBER = '#,##0.0;(#,##0.0);"—"'
PERCENT = '0.0%;(0.0%)'
MULTIPLE = '0.0"x"'
LIGHT_BORDER = Side(style="thin", color="D6DEE6")
BASE_BORDER = Border(**{k: Side(style="thin", color="1676D2") for k in ("left", "right", "top", "bottom")})


def write(ws, row, col, value, fmt=NUMBER, *, input_=False, warning=False, bold=False):
    cell = ws.cell(row, col, value)
    # User-supplied context is text, never executable workbook formulas.
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Arial", size=10, color=RED if warning else BLUE if input_ else "202B38", bold=bold)
    cell.number_format = fmt
    cell.alignment = Alignment(horizontal="right" if isinstance(value, (int, float)) or value == "N/A" else "left", vertical="center")
    return cell


def section(ws, row, title, columns=7):
    for col in range(1, columns + 1):
        cell = write(ws, row, col, title if col == 1 else None, bold=True)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    ws.row_dimensions[row].height = 23
    if title != "LINE ITEM":
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)


def note(ws, row, text, columns=7, warning=False):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = write(ws, row, 1, text, warning=warning)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30


def heat(irr):
    return "E9EDF1" if irr is None else "F5DEDE" if irr < 0.15 else "FFF0C2" if irr <= 0.2 else "DDEFE3"
