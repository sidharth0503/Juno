"""Workbook layout only: all financial values come from the existing engines."""
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.financial import run_lbo_model
from app.financial.sensitivity import run_sensitivity_analysis
from app.schemas.model import ExportRequest
from app.export.styles import BASE_BORDER, BLUE, DARK, LIGHT_BORDER, MULTIPLE, NUMBER, PERCENT, heat, note, section, write

MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEETS = ("LBO Summary", "Operating Model", "Income Statement", "Cash Flow", "Balance Sheet", "Debt Schedule", "Sensitivity")
SUMMARY_CELLS = {"initial_equity": "B29", "exit_ev": "B34", "exit_equity": "B36", "moic": "B37", "irr": "B38"}


def export_filename(request: ExportRequest) -> str:
    ticker = request.company.ticker if request.company else None
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", ticker or "").strip("._-")[:32]
    return f"{safe}_LBO_Model.xlsx" if safe else "LBO_Model.xlsx"


def build_excel(request: ExportRequest) -> bytes:
    # Recalculate from validated inputs. No client outputs are accepted.
    result = run_lbo_model(request.historical, request.transaction_assumptions, request.operating_assumptions)
    sensitivity = run_sensitivity_analysis(request.historical, request.transaction_assumptions,
        request.operating_assumptions, request.entry_multiples, request.exit_multiples)
    wb = Workbook()
    wb.remove(wb.active)
    company = request.company
    currency = company.currency if company else None
    context = f"{company.company_name or company.ticker or 'Company model' if company else 'Manual model'} — {currency or 'Currency unspecified'} millions"

    def sheet(name, columns=7):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C6"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:5"
        ws.column_dimensions["A"].width = 34
        for col in range(2, columns + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        ws.sheet_format.defaultRowHeight = 20
        ws.sheet_properties.tabColor = DARK
        note(ws, 1, name.upper(), columns)
        ws["A1"].font = Font(name="Arial", size=14, bold=True, color=DARK)
        note(ws, 2, context, columns)
        note(ws, 3, "Python model snapshot. Editing workbook inputs does not recalculate outputs; re-export from the terminal.", columns)
        return ws

    def years(ws):
        section(ws, 5, "LINE ITEM")
        for col, title in enumerate([f"Base FY{company.fiscal_year}" if company and company.fiscal_year else "Base year", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"], 2):
            ws.cell(5, col).value = title
            ws.cell(5, col).alignment = Alignment(horizontal="center")

    def row(ws, row_number, label, field=None, base=None, fmt=NUMBER, assumption=None, warning=False, base_input=True):
        write(ws, row_number, 1, label, bold=warning)
        write(ws, row_number, 2, base, fmt, input_=base is not None and base_input)
        for col, year in enumerate(result.projected_years, 3):
            value = getattr(year, field) if field else assumption
            cell = write(ws, row_number, col, value, fmt, input_=field is None, warning=warning and value != 0)
            if warning:
                cell.border = Border(top=LIGHT_BORDER)

    ws = sheet("LBO Summary")
    ws.freeze_panes = "B5"
    for r, label, value in [(5, "Company", company.company_name if company else "Manual model"),
        (6, "Ticker", company.ticker if company else None), (7, "Currency", currency or "Unspecified"),
        (8, "Base fiscal year", company.fiscal_year if company else None),
        (9, "Data source", company.provider if company else "Manual inputs")]:
        write(ws, r, 1, label)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        write(ws, r, 2, value, "0" if r == 8 else "General", input_=True)
    t, o = request.transaction_assumptions, request.operating_assumptions
    section(ws, 11, "TRANSACTION ASSUMPTIONS")
    for r, label, value, fmt in [(12,"Purchase multiple",t.purchase_multiple_ev_ebitda,MULTIPLE),
        (13,"Debt financing",t.debt_percentage,PERCENT),(14,"Interest rate",t.interest_rate,PERCENT),(15,"Exit multiple",t.exit_multiple,MULTIPLE)]:
        write(ws,r,1,label); write(ws,r,2,value,fmt,input_=True)
    section(ws,17,"OPERATING ASSUMPTIONS")
    for r,label,value,fmt in [(18,"Revenue growth",o.revenue_growth_rate,PERCENT),(19,"EBITDA margin",o.ebitda_margin,PERCENT),
        (20,"Tax rate",o.tax_rate,PERCENT),(21,"DSO (days)",o.dso_days,"0"),(22,"DIO (days)",o.dio_days,"0"),
        (23,"DPO (days)",o.dpo_days,"0"),(24,"CapEx / revenue",o.capex_pct_of_rev,PERCENT)]:
        write(ws,r,1,label); write(ws,r,2,value,fmt,input_=True)
    section(ws,26,"TRANSACTION SUMMARY")
    for r,label,value in [(27,"Purchase enterprise value",result.transaction.purchase_enterprise_value),
        (28,"Initial debt",result.transaction.initial_debt),(29,"Initial equity",result.transaction.initial_equity)]:
        write(ws,r,1,label); write(ws,r,2,value,bold=True)
    section(ws,31,"EXIT SUMMARY")
    for r,label,value,fmt in [(32,"Year 5 EBITDA",result.projected_years[-1].ebitda,NUMBER),(33,"Exit multiple",t.exit_multiple,MULTIPLE),
        (34,"Exit enterprise value",result.returns.exit_enterprise_value,NUMBER),(35,"Exit debt",result.returns.exit_debt,NUMBER),
        (36,"Exit equity",result.returns.exit_equity,NUMBER),(37,"MOIC",result.returns.moic,MULTIPLE),(38,"IRR",result.returns.irr,PERCENT)]:
        write(ws,r,1,label); write(ws,r,2,"N/A" if value is None else value,fmt,bold=True,warning=value is None)
    section(ws,40,"CURRENT BASE INPUTS (INCLUDING MANUAL OVERRIDES)")
    historical_labels = [("revenue","Revenue"),("cogs","COGS"),("ebitda","EBITDA"),("cash","Cash"),("accounts_receivable","Accounts receivable"),
        ("inventory","Inventory"),("accounts_payable","Accounts payable"),("ppe","Net PP&E"),("debt","Debt"),("shareholders_equity","Shareholders equity")]
    for r,(key,label) in enumerate(historical_labels,41):
        write(ws,r,1,label); write(ws,r,2,getattr(request.historical,key),input_=True)
    section(ws,52,"MODEL NOTES AND DIAGNOSTICS")
    notes = ["Amounts are in millions. Blue numbers are inputs; dark numbers are Python outputs.",
        "IRR is N/A when exit equity is nonpositive. Blank historical cells on schedules mean unavailable/not applicable.",
        "Provider context describes the imported company; current base inputs may contain manual overrides."]
    for r,text in enumerate(notes,53): note(ws,r,text)
    for r,d in enumerate(result.diagnostics,56): note(ws,r,f"Year {d.year}: {d.message}",warning=True)

    ws = sheet("Operating Model"); years(ws)
    section(ws,6,"OPERATING PROJECTION")
    h=result.historical
    row(ws,7,"Revenue","revenue",h.revenue)
    row(ws,8,"Revenue growth",fmt=PERCENT,assumption=o.revenue_growth_rate)
    row(ws,9,"COGS","cogs",h.cogs)
    row(ws,10,"EBITDA","ebitda",h.ebitda)
    row(ws,11,"EBITDA margin",fmt=PERCENT,assumption=o.ebitda_margin)
    section(ws,13,"WORKING CAPITAL")
    for r,label,key in [(14,"Accounts receivable","accounts_receivable"),(15,"Inventory","inventory"),(16,"Accounts payable","accounts_payable")]:
        row(ws,r,label,key,getattr(h,key))
    row(ws,17,"Net working capital","nwc",result.projected_years[0].prior_nwc,base_input=False)
    row(ws,18,"Change in net working capital","change_in_nwc")
    section(ws,20,"PROPERTY, PLANT AND EQUIPMENT")
    for r,label,key,base in [(21,"Beginning PP&E","beginning_ppe",None),(22,"CapEx","capex",None),
        (23,"Depreciation","depreciation",None),(24,"Ending PP&E","ending_ppe",h.ppe)]: row(ws,r,label,key,base)

    ws=sheet("Income Statement"); years(ws)
    for r,label,key,base in [(6,"Revenue","revenue",h.revenue),(7,"COGS","cogs",h.cogs),(8,"EBITDA","ebitda",h.ebitda),
        (9,"D&A (modeled depreciation)","depreciation",None),(10,"EBIT","ebit",None),(11,"Interest expense","interest_expense",None),
        (12,"EBT","ebt",None),(13,"Taxes","taxes",None),(14,"Net income","net_income",None)]: row(ws,r,label,key,base)
    note(ws,16,"Expense and paydown rows retain the engine's positive expense convention; no additional sign or scale conversion.")

    ws=sheet("Cash Flow"); years(ws)
    for r,(label,key) in enumerate([("Net income","net_income"),("Depreciation","depreciation"),("Change in NWC","change_in_nwc"),
        ("CapEx","capex"),("CFADS","cfads"),("Mandatory debt paydown","mandatory_debt_paydown"),("Optional debt paydown","optional_debt_paydown"),
        ("Retained cash flow","retained_cash_flow"),("Ending cash","ending_cash")],6): row(ws,r,label,key,h.cash if key=="ending_cash" else None)

    ws=sheet("Balance Sheet"); years(ws)
    for r,label,key,base in [(6,"Cash","ending_cash",h.cash),(7,"Accounts receivable","accounts_receivable",h.accounts_receivable),
        (8,"Inventory","inventory",h.inventory),(9,"Net PP&E","ending_ppe",h.ppe),(10,"Total assets","total_assets",None),
        (12,"Accounts payable","accounts_payable",h.accounts_payable),(13,"Debt","ending_debt",h.debt),
        (14,"Total liabilities","total_liabilities",None),(16,"Equity","ending_equity",h.shareholders_equity),
        (18,"Balance check","balance_check",None)]: row(ws,r,label,key,base,warning=key=="balance_check")
    note(ws,20,"Simplified transaction balance sheet: acquisition accounting is not modeled. Imbalances are retained, not plugged.",warning=True)

    ws=sheet("Debt Schedule"); years(ws)
    row(ws,6,"Beginning debt","beginning_debt")
    row(ws,7,"Interest rate",fmt=PERCENT,assumption=t.interest_rate)
    for r,label,key in [(8,"Interest expense","interest_expense"),(9,"Mandatory debt paydown","mandatory_debt_paydown"),
        (10,"Optional debt paydown","optional_debt_paydown"),(11,"Ending debt","ending_debt")]: row(ws,r,label,key,h.debt if key=="ending_debt" else None)
    note(ws,13,"Single acquisition debt balance. Historical debt is reference-only and is replaced by acquisition debt.")

    columns=max(7,len(sensitivity.exit_multiples)+1)
    ws=sheet("Sensitivity",columns); ws.freeze_panes="B6"
    section(ws,4,"IRR: ROWS = ENTRY MULTIPLE; COLUMNS = EXIT MULTIPLE",columns)
    write(ws,5,1,"Entry / Exit",bold=True)
    for c,value in enumerate(sensitivity.exit_multiples,2): write(ws,5,c,value,MULTIPLE,input_=True,bold=True)
    for r,cells in enumerate(sensitivity.matrix,6):
        write(ws,r,1,sensitivity.entry_multiples[r-6],MULTIPLE,input_=True,bold=True)
        for c,item in enumerate(cells,2):
            cell=write(ws,r,c,"N/A" if item.irr is None else item.irr,PERCENT)
            cell.fill=PatternFill("solid",fgColor=heat(item.irr))
            if (r-6,c-2)==(sensitivity.base_case.row_index,sensitivity.base_case.column_index): cell.border=BASE_BORDER
    end=7+len(sensitivity.matrix)
    note(ws,end,"IRR: red <15%; amber 15–20%; green >20%; gray N/A. Blue outline marks the current base case.",columns)
    note(ws,end+1,f"Base assumptions: entry {t.purchase_multiple_ev_ebitda:g}x, exit {t.exit_multiple:g}x. Each cell runs the full Python model.",columns)
    for r,warning in enumerate(sensitivity.warnings,end+2): note(ws,r,warning,columns,warning=True)

    for ws in wb:
        ws.print_area=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    output=BytesIO()
    wb.save(output)
    return output.getvalue()
