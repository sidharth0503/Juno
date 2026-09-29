from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.export.excel import MIME_TYPE, SHEETS, build_excel, export_filename
from app.export.styles import BLUE, heat
from app.schemas.model import ExportRequest
from app.financial import run_lbo_model
from app.financial.sensitivity import run_sensitivity_analysis
from test_model_api import payload


def workbook(payload):
    return load_workbook(BytesIO(build_excel(ExportRequest.model_validate(payload))))


def test_golden_values_and_structure(payload):
    request = ExportRequest.model_validate(payload)
    wb = workbook(payload)
    result = run_lbo_model(request.historical, request.transaction_assumptions, request.operating_assumptions)
    assert wb.sheetnames == list(SHEETS)
    expected = {"B29": result.transaction.initial_equity, "B34": result.returns.exit_enterprise_value,
        "B36": result.returns.exit_equity, "B37": result.returns.moic, "B38": result.returns.irr}
    for address, value in expected.items():
        cell = wb["LBO Summary"][address]
        assert cell.value == pytest.approx(value)
        assert cell.data_type == "n"
    assert wb["Income Statement"]["C6"].value == 1050
    assert wb["Income Statement"]["G8"].value == pytest.approx(result.projected_years[-1].ebitda)
    assert wb["Debt Schedule"]["G11"].value == pytest.approx(result.returns.exit_debt)
    assert wb["Sensitivity"]["D8"].value == pytest.approx(result.returns.irr)
    assert wb["Balance Sheet"]["C18"].value == pytest.approx(-1900)
    assert "Simplified" in wb["Balance Sheet"]["A20"].value
    assert wb["LBO Summary"]["B6"].value is None
    for ws in wb:
        assert ws.freeze_panes
        assert ws.sheet_view.showGridLines is False
        assert ws.print_area
        assert not any(c.data_type == "f" for row in ws for c in row)
    assert not wb._external_links


def test_workbook_styles_and_types(payload):
    wb = workbook(payload)
    summary = wb["LBO Summary"]
    assert summary["B12"].font.color.rgb.endswith(BLUE)
    assert summary["B29"].font.color.rgb.endswith("202B38")
    assert summary["A11"].fill.fgColor.rgb.endswith("172B42")
    assert "%" in summary["B38"].number_format
    assert '"x"' in summary["B37"].number_format
    assert summary["B38"].data_type == "n"
    base = wb["Sensitivity"]["D8"]
    assert base.border.left.style == "thin"
    assert base.border.left.color.rgb.endswith("1676D2")
    assert base.fill.fgColor.rgb.endswith(heat(base.value))


def test_export_null_irr(payload):
    payload["operating_assumptions"]["ebitda_margin"] = 0
    wb = workbook(payload)
    assert wb["LBO Summary"]["B38"].value == "N/A"
    assert wb["Sensitivity"]["D8"].value == "N/A"
    assert wb["LBO Summary"]["B37"].data_type == "n"


def test_export_preserves_millions_and_manual_override(payload):
    payload["historical"]["revenue"] = 416161
    payload["historical"]["ebitda"] = 150000
    payload["company"] = {"ticker": "AAPL", "company_name": "Apple Inc.", "currency": "USD", "fiscal_year": 2025, "provider": "yfinance"}
    wb = workbook(payload)
    assert wb["LBO Summary"]["B41"].value == 416161
    assert wb["LBO Summary"]["B43"].value == 150000
    assert wb["Operating Model"]["B7"].value == 416161
    assert wb["Operating Model"]["B10"].value == 150000
    assert wb["LBO Summary"]["B6"].value == "AAPL"
    assert wb["LBO Summary"]["B8"].value == 2025


def test_custom_grid_orientation_and_base(payload):
    payload.update(entry_multiples=[12, 9], exit_multiples=[11, 8, 10])
    payload["transaction_assumptions"].update(purchase_multiple_ev_ebitda=9, exit_multiple=11)
    request = ExportRequest.model_validate(payload)
    result = run_sensitivity_analysis(request.historical, request.transaction_assumptions, request.operating_assumptions, request.entry_multiples, request.exit_multiples)
    ws = workbook(payload)["Sensitivity"]
    assert ws["A6"].value == 12 and ws["A7"].value == 9
    assert [ws.cell(5,c).value for c in range(2,5)] == [11,8,10]
    for r,row in enumerate(result.matrix,6):
        for c,cell in enumerate(row,2):
            assert ws.cell(r,c).value == pytest.approx(cell.irr)
    assert ws["B7"].border.left.color.rgb.endswith("1676D2")
    assert ws["D8"].value is None


def test_financial_values_deterministic(payload):
    first, second = workbook(payload), workbook(payload)
    for name in SHEETS:
        assert list(first[name].values) == list(second[name].values)


def test_http_xlsx_and_cors(payload):
    response = TestClient(app).post("/api/v1/model/export", json=payload, headers={"Origin": "http://localhost:3000"})
    assert response.status_code == 200
    assert response.headers["content-type"] == MIME_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="LBO_Model.xlsx"'
    assert "Content-Disposition" in response.headers["access-control-expose-headers"]
    wb = load_workbook(BytesIO(response.content))
    assert wb["LBO Summary"]["B29"].value == pytest.approx(900)


@pytest.mark.parametrize("change,status", [("missing",422),("domain",400),("outputs",422),("grid",422)])
def test_invalid_export_is_not_a_file(payload, change, status):
    if change == "missing": payload["historical"].pop("cogs")
    elif change == "domain": payload["historical"]["ebitda"] = 0
    elif change == "outputs": payload["result"] = {"irr": 999}
    else: payload["entry_multiples"] = [1] * 11
    response = TestClient(app).post("/api/v1/model/export", json=payload)
    assert response.status_code == status
    assert "application/json" in response.headers["content-type"]


def test_metadata_is_literal_and_filename_sanitized(payload):
    payload["company"] = {"ticker": "../../BAD\\NAME", "company_name": "=1+1", "currency": "EUR"}
    request = ExportRequest.model_validate(payload)
    assert "/" not in export_filename(request) and "\\" not in export_filename(request)
    ws = workbook(payload)["LBO Summary"]
    assert ws["B5"].value == "=1+1"
    assert ws["B5"].data_type == "s"
    assert "EUR millions" in ws["A2"].value


@pytest.mark.parametrize("irr,color", [(None,"E9EDF1"),(0.149,"F5DEDE"),(0.15,"FFF0C2"),(0.2,"FFF0C2"),(0.201,"DDEFE3")])
def test_sensitivity_style_thresholds(irr,color):
    assert heat(irr) == color
